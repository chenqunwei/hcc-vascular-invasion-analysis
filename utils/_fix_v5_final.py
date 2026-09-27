# -*- coding: utf-8 -*-
"""
Final fix for v5_fixed.docx (Plan A):
 1) text replacements (DFS/OS log-rank P, PAC, ΔCDF, Table-1 three stale cells)
 2) re-embed Fig2-5 (rId7-10) with the authoritative results/ PNGs (same pixel dims)
 3) afterwards: restore results/Fig1 from figures/Fig1 and sync figures/Fig2-5 <- results/
Backs up v5_fixed.docx before overwriting.
"""
import os, zipfile, shutil, xml.etree.ElementTree as ET
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table

SRC = "HCC血管侵犯连续谱_手稿_v5_fixed.docx"
BAK = SRC + ".bak2"
RELS_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"

EN = "\u2013"

# ---- text replacement map (PRE -> authoritative) ----
REPL = [
    ("PAC = 0.137", "PAC = 0.139"),
    ("0.176", "0.175"),                              # ΔCDF area
    ("log-rank P = 0.006", "log-rank P = 0.013"),    # DFS (L51, L73, + inside L66)
    ("OS P = 0.194", "OS P = 0.265"),                # L66
    ("log-rank P = 0.194", "log-rank P = 0.265"),    # L51, L73
    ("36 (50.7%)", "37 (51.4%)"),                    # Table1 Male High VI
    ("19 (28.4%)", "19 (27.9%)"),                    # Table1 Stage III/IV High VI
    ("40 (90.9%)", "41 (91.1%)"),                    # Table1 Child-Pugh A High VI
]

# rId -> authoritative new PNG (same pixel dimensions as embedded -> lossless swap)
NEW_IMG = {
    "rId7":  "results/Fig2_molecular_subtyping.png",
    "rId8":  "results/Fig3_MVI_clinical.png",
    "rId9":  "results/Fig4_survival.png",
    "rId10": "results/Fig5_immune.png",
}

# ----------------------------------------------------------------------
def replace_one(runs, a, b, new):
    pos = 0; spans = []
    for r in runs:
        s = pos; e = pos + len(r.text); spans.append((s, e, r)); pos = e
    new_texts = {}
    for (s, e, r) in spans:
        L = len(r.text)
        head_end = max(0, min(e, a) - s)
        tail_start = max(s, b) - s
        head = r.text[:head_end]; tail = r.text[tail_start:] if tail_start <= L else ""
        if s <= a < e:
            new_texts[r] = head + new + tail
        else:
            new_texts[r] = head + tail
    for r, nt in new_texts.items():
        r.text = nt

def replace_all_in_runs(runs, old, new, max_iter=200):
    cnt = 0
    for _ in range(max_iter):
        full = "".join(r.text for r in runs)
        a = full.find(old)
        if a < 0: break
        replace_one(runs, a, a + len(old), new); cnt += 1
    return cnt

def iter_block_items(parent):
    from docx.oxml.ns import qn
    for child in parent.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

def block_runs(blk):
    if isinstance(blk, Paragraph):
        return [blk.runs]
    containers = []
    for row in blk.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                containers.append(pp.runs)
    return containers

def main():
    # backup
    if not os.path.exists(BAK):
        shutil.copy(SRC, BAK)
        print(f"[BAK] saved {BAK}")

    doc = docx.Document(SRC)
    total = 0
    for blk in iter_block_items(doc):
        for runs in block_runs(blk):
            if not runs: continue
            for old, new in REPL:
                total += replace_all_in_runs(runs, old, new)
    doc.save(SRC)
    print(f"[TEXT] applied {total} targeted edits -> {SRC}")

    # ---- re-embed Fig2-5 via zip byte swap (dynamic rId -> target) ----
    with zipfile.ZipFile(SRC) as z:
        rels = z.read("word/_rels/document.xml.rels")
        names = set(z.namelist())
    root = ET.fromstring(rels)
    rid_target = {r.get("Id"): r.get("Target") for r in root.findall(f"{RELS_NS}Relationship")}
    swap = {}
    for rid, png in NEW_IMG.items():
        tgt = rid_target.get(rid, "")
        if not tgt:
            print(f"[WARN] {rid} has no Target"); continue
        entry = "word/" + tgt if not tgt.startswith("/") else tgt[1:]
        if entry not in names:
            print(f"[WARN] {entry} not in zip (rid={rid})"); continue
        swap[entry] = png
        print(f"[IMG] {rid} -> {entry} <- {png}")

    tmp = SRC + ".tmp"
    with zipfile.ZipFile(SRC) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in swap:
                data = open(swap[item.filename], "rb").read()
                print(f"[IMG] replaced {item.filename} ({len(data)} bytes)")
            zout.writestr(item, data)
    os.replace(tmp, SRC)
    print(f"[ZIP] re-embedded Fig2-5, wrote {SRC}")

    # ---- restore / sync figure files ----
    if os.path.exists("figures/Fig1_vascular_program.png"):
        shutil.copy("figures/Fig1_vascular_program.png", "results/Fig1_vascular_program.png")
        print("[SYNC] restored results/Fig1 <- figures/Fig1 (correct portrait)")
    for f in ["Fig2_molecular_subtyping","Fig3_MVI_clinical","Fig4_survival","Fig5_immune"]:
        s, d = f"results/{f}.png", f"figures/{f}.png"
        if os.path.exists(s):
            shutil.copy(s, d); print(f"[SYNC] figures/{f}.png <- results/{f}.png")

    # ---- verify text residuals ----
    doc2 = docx.Document(SRC)
    txt = []
    for blk in iter_block_items(doc2):
        if isinstance(blk, Paragraph):
            if blk.text.strip(): txt.append(blk.text)
        else:
            for row in blk.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip(): txt.append(p.text)
    full = "\n".join(txt)
    pairs = [("PAC = 0.137","PAC = 0.139"),("0.176","0.175"),
             ("log-rank P = 0.006","log-rank P = 0.013"),("OS P = 0.194","OS P = 0.265"),
             ("log-rank P = 0.194","log-rank P = 0.265"),
             ("36 (50.7%)","37 (51.4%)"),("19 (28.4%)","19 (27.9%)"),("40 (90.9%)","41 (91.1%)")]
    resid = [o for o,n in pairs if o in full]
    print("\n[VERIFY] residual OLD tokens:", resid if resid else "NONE")
    ok = True
    for o,n in pairs:
        present = (n in full) and (o not in full)
        ok = ok and present
        print(f"   {'OK ' if present else 'FAIL'} {o!r} -> {n!r}")
    print("\nRESULT:", "PASS" if not resid else "FAIL (old values remain)")

if __name__ == "__main__":
    main()
