# -*- coding: utf-8 -*-
"""
Fix v5.docx: replace all PRE-fix (seed-None) numbers with the NEW deterministic
(post-fix, seed=42 + VI-order relabel) values.

The previous attempt failed because it used run-level `run.text.replace` on whole
multi-token strings; when a sentence is split across several <w:r> runs (bold,
reference markers, etc.) the exact string never sits inside a single run, so the
replacement silently did nothing.

This script uses a paragraph-level replacement that is robust to text being split
across runs: it concatenates all run texts, finds the (possibly cross-run) match,
and rewrites the affected runs while preserving per-run formatting. It also rewrites
table cells the same way.

Because Word currently has v5.docx open (lock file present), the result is written
to a NEW file `HCC..._手稿_v5_fixed.docx` to avoid fighting the lock.
"""
import os
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table

EN = "\u2013"  # en dash used in CIs

SRC = "HCC血管侵犯连续谱_手稿_v5.docx"
DST = "HCC血管侵犯连续谱_手稿_v5_fixed.docx"

# ----------------------------------------------------------------------
# Robust cross-run replacement
# ----------------------------------------------------------------------
def replace_one(runs, a, b, new):
    """Replace full-text span [a,b) with `new`, distributing across runs."""
    pos = 0
    spans = []
    for r in runs:
        s = pos
        e = pos + len(r.text)
        spans.append((s, e, r))
        pos = e
    new_texts = {}
    for (s, e, r) in spans:
        L = len(r.text)
        head_end = max(0, min(e, a) - s)      # chars kept before the match
        tail_start = max(s, b) - s            # chars kept after the match
        head = r.text[:head_end]
        tail = r.text[tail_start:] if tail_start <= L else ""
        if s <= a < e:                        # run that contains match start
            new_texts[r] = head + new + tail
        else:
            new_texts[r] = head + tail
    for r, nt in new_texts.items():
        r.text = nt

def replace_all_in_runs(runs, old, new, max_iter=200):
    cnt = 0
    for _ in range(max_iter):
        full = "".join(r.text for r in runs)
        a = full.find(old)
        if a < 0:
            break
        replace_one(runs, a, a + len(old), new)
        cnt += 1
    return cnt

# ----------------------------------------------------------------------
# Replacement map: PRE-fix -> NEW deterministic
# (ordered: specific/long strings BEFORE short substrings to avoid partial hits)
# ----------------------------------------------------------------------
REPL = [
    # ---- body: subtyping sentence ----
    ("Low VI (n = 100), Intermediate VI (n = 128), and High VI (n = 71)",
     "Low VI (n = 99), Intermediate VI (n = 128), and High VI (n = 72)"),
    ("(499, 705, and 914, respectively)", "(497, 705, and 914, respectively)"),
    ("21.0% (Low), 32.8% (Intermediate), and 42.3% (High VI)",
     "21.2% (Low), 32.8% (Intermediate), and 41.7% (High VI)"),
    # ---- MVI chi-square (body + Fig3 legend) ----
    ("9.06, P = 0.011", "8.44, P = 0.0147"),
    # ---- AFP ----
    ("median 1,412 ng/mL", "median 1,456 ng/mL"),
    ("Low (12.5 ng/mL)", "Low (12.0 ng/mL)"),
    ("median 1,412 vs 12.5 ng/mL", "median 1,456 vs 12.0 ng/mL"),
    # ---- stage / sex clinical sentences ----
    ("28.4% vs 14.1% and 17.2%, P = 0.064", "27.9% vs 14.3% and 17.2%, P = 0.077"),
    ("50.7% vs 68.0% and 74.2%, P = 0.003", "51.4% vs 67.7% and 74.2%, P = 0.004"),
    # ---- Cox High vs Low (body) ----
    ("1.97-fold higher recurrence risk", "1.82-fold higher recurrence risk"),
    ("HR = 1.97, 95% CI 1.16" + EN + "3.33, P = 0.012",
     "HR = 1.82, 95% CI 1.08" + EN + "3.05, P = 0.024"),
    # ---- Cox per level (body + discussion) ----
    ("HR = 1.43 per subtype level, 95% CI 1.12" + EN + "1.82, P = 0.0043",
     "HR = 1.38 per subtype level, 95% CI 1.08" + EN + "1.76, P = 0.0091"),
    ("rose monotonically from 21% in the Low VI subtype to 42% in the High VI subtype, "
     "and recurrence risk increased by 43% per subtype level (HR = 1.43, P = 0.0043)",
     "rose monotonically from 21.2% in the Low VI subtype to 41.7% in the High VI subtype, "
     "and recurrence risk increased by 38% per subtype level (HR = 1.38, P = 0.0091)"),
    # ---- Table 1 header ----
    ("Low VI (n=100)", "Low VI (n=99)"),
    ("High VI (n=71)", "High VI (n=72)"),
    # ---- Table 1 rows ----
    ("59.5", "60.0"),
    ("68 (68.0%)", "67 (67.7%)"),
    ("13 (14.1%)", "13 (14.3%)"),
    ("21 (21.0%)", "21 (21.2%)"),
    ("12.5", "12.0"),
    ("1412.0", "1456.0"),
    ("65 (94.2%)", "64 (94.1%)"),
    # ---- Table 2 rows (Cox) ----
    ("1.97 (1.16" + EN + "3.33)", "1.82 (1.08" + EN + "3.05)"),
    ("1.41 (0.79" + EN + "2.51)", "1.34 (0.76" + EN + "2.38)"),
    ("0.244", "0.310"),
    ("1.43 (1.12" + EN + "1.82)", "1.38 (1.08" + EN + "1.76)"),
    ("0.0043", "0.0091"),
    ("1.16 (0.87" + EN + "1.55)", "1.14 (0.85" + EN + "1.52)"),
    ("0.318", "0.383"),
    # ---- Abstract (bare MVI prevalence, no labels) ----
    ("21.0%, 32.8%, and 42.3%", "21.2%, 32.8%, and 41.7%"),
    # ---- Table 1 bare P-value / percentage cells (not caught above) ----
    ("30 (42.3%)", "30 (41.7%)"),     # MVI High %
    ("0.067", "0.065"),               # Age P
    ("0.003", "0.004"),               # Male P
    ("0.064", "0.077"),               # Stage P
    ("0.011", "0.0147"),              # MVI P (body/legend handled earlier)
    ("0.464", "0.477"),               # Child-Pugh P
    # ---- Table 2 bare P-value cell ----
    ("0.012", "0.024"),               # High vs Low DFS P (body handled earlier)
]

# Tokens that must be GONE after the fix (residual check)
OLD_TOKENS = [
    "n = 100", "n = 71", "(499, 705", "21.0% (Low)", "42.3% (High",
    "9.06, P = 0.011", "median 1,412", "12.5 ng/mL", "28.4% vs 14.1%",
    "50.7% vs 68.0%", "1.97-fold", "1.97, 95%", "1.43 per subtype",
    "1.43 (1.12", "59.5", "68 (68.0%)", "13 (14.1%)", "21 (21.0%)",
    "1412.0", "65 (94.2%)", "1.97 (1.16", "1.41 (0.79", "0.244",
    "1.16 (0.87", "0.318",
    # broadened to catch bare table/abstract values
    "21.0%", "42.3%", "21.0%, 32.8%, and 42.3%", "30 (42.3%)",
    "0.067", "0.003", "0.064", "0.464", "0.011", "0.012",
]
# Tokens that must be PRESENT after the fix
NEW_TOKENS = [
    "n = 99", "n = 72", "(497, 705", "21.2% (Low)", "41.7% (High",
    "8.44, P = 0.0147", "median 1,456", "12.0 ng/mL", "27.9% vs 14.3%",
    "51.4% vs 67.7%", "1.82-fold", "1.82, 95%", "1.38 per subtype",
    "1.38 (1.08", "60.0", "67 (67.7%)", "13 (14.3%)", "21 (21.2%)",
    "1456.0", "64 (94.1%)", "1.82 (1.08", "1.34 (0.76", "0.310",
    "1.14 (0.85", "0.383",
    # broadened
    "21.2%, 32.8%, and 41.7%", "30 (41.7%)",
    "0.065", "0.004", "0.077", "0.477", "0.024",
]

# ----------------------------------------------------------------------
def iter_block_items(parent):
    from docx.oxml.ns import qn
    for child in parent.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

def block_runs(blk):
    if isinstance(blk, Paragraph):
        return [blk.runs]
    # Each table CELL is an independent container so replacements never
    # bridge cell boundaries (which previously corrupted adjacent cells).
    containers = []
    for row in blk.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                containers.append(pp.runs)
    return containers

def main():
    doc = docx.Document(SRC)
    total = 0
    for blk in iter_block_items(doc):
        containers = block_runs(blk)
        for runs in containers:
            if not runs:
                continue
            for old, new in REPL:
                c = replace_all_in_runs(runs, old, new)
                if c:
                    total += c
    doc.save(DST)

    # ---- verification ----
    doc2 = docx.Document(DST)
    txt = []
    for blk in iter_block_items(doc2):
        if isinstance(blk, Paragraph):
            t = blk.text
            if t:
                txt.append(t)
        else:
            for row in blk.rows:
                for cell in row.cells:
                    txt.append(cell.text)
    full = "\n".join(txt)
    residual = [t for t in OLD_TOKENS if t in full]
    present = [t for t in NEW_TOKENS if t in full]
    print(f"[OK] wrote {DST}")
    print(f"[OK] total targeted edits applied: {total}")
    print(f"[VERIFY] residual OLD tokens still present: {len(residual)}")
    for t in residual:
        print("   RESIDUAL:", repr(t))
    print(f"[VERIFY] NEW tokens present: {len(present)}/{len(NEW_TOKENS)}")
    missing = [t for t in NEW_TOKENS if t not in full]
    for t in missing:
        print("   MISSING-NEW:", repr(t))
    if residual:
        print("RESULT: FAIL (old values remain)")
    else:
        print("RESULT: PASS (no residual old values)")

if __name__ == "__main__":
    main()
