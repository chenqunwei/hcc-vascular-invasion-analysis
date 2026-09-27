# -*- coding: utf-8 -*-
"""将手稿 v3 (markdown) 转换为 Word (.docx)。"""
import re
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SRC = "HCC血管侵犯连续谱_手稿_v3.md"
OUT = "HCC血管侵犯连续谱_手稿_v3.docx"

TOKEN = re.compile(
    r'(\*\*[^*]+\*\*'      # **bold**
    r'|\*[^*]+\*'          # *italic*
    r'|<sub>.*?</sub>'     # <sub>下标</sub>
    r'|<sup>.*?</sup>'     # <sup>上标</sup>
    r'|`[^`]+`'            # `code`
    r'|\[\d+\])'           # [n] 引用 → 上标
)


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_md_runs(paragraph, text, size=12):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            set_run_font(paragraph.add_run(text[pos:m.start()]), size)
        token = m.group(0)
        if token.startswith('**') and token.endswith('**'):
            set_run_font(paragraph.add_run(token[2:-2]), size, bold=True)
        elif token.startswith('*') and token.endswith('*'):
            set_run_font(paragraph.add_run(token[1:-1]), size, italic=True)
        elif token.startswith('<sub>'):
            r = paragraph.add_run(token[5:-6]); set_run_font(r, size); r.font.subscript = True
        elif token.startswith('<sup>'):
            r = paragraph.add_run(token[5:-6]); set_run_font(r, size); r.font.superscript = True
        elif token.startswith('`'):
            r = paragraph.add_run(token[1:-1]); set_run_font(r, size); r.font.name = 'Consolas'
        elif token.startswith('['):
            r = paragraph.add_run(token); set_run_font(r, size); r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size)


def add_heading(doc, text, level):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    return p


def add_body(doc, text, size=12, line_spacing=1.5, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    add_md_runs(p, text, size)
    return p


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    add_md_runs(p, text, size=11)
    return p


def parse_table(lines, start_idx):
    rows = []
    i = start_idx
    header = [c.strip() for c in lines[i].strip().strip('|').split('|')]
    rows.append(header)
    i += 1
    if i < len(lines) and re.match(r'^\s*\|[\s:\-|]+\|\s*$', lines[i]):
        i += 1
    while i < len(lines) and lines[i].strip().startswith('|'):
        rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
        i += 1
    return rows, i


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            p = table.cell(r_idx, c_idx).paragraphs[0]
            add_md_runs(p, cell_text, size=10)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0
                               else WD_ALIGN_PARAGRAPH.CENTER)
    return table


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    lines = open(SRC, encoding='utf-8').read().split('\n')
    section = 'body'  # body | refs | figures

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped or stripped == '---' or stripped.startswith('>'):
            i += 1
            continue

        # 图片
        img_m = re.match(r'^!\[.*?\]\((.+?)\)$', stripped)
        if img_m:
            path = img_m.group(1)
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(6.0))
            i += 1
            continue

        # 主标题 #
        h1 = re.match(r'^#\s+(.*)$', stripped)
        if h1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(h1.group(1).strip()), 16, bold=True)
            p.paragraph_format.space_after = Pt(12)
            i += 1
            continue

        # 二级标题 ##
        h2 = re.match(r'^##\s+(.*)$', stripped)
        if h2:
            text = h2.group(1).strip()
            if text.startswith('Figures（附图'):
                add_heading(doc, 'Figures', 1)
                section = 'figures'
            elif text.startswith('References'):
                add_heading(doc, 'References', 1)
                section = 'refs'
            else:
                add_heading(doc, text, 1)
                section = 'body'
            i += 1
            continue

        # 三级标题 ###
        h3 = re.match(r'^###\s+(.*)$', stripped)
        if h3:
            add_heading(doc, h3.group(1).strip(), 2)
            i += 1
            continue

        # markdown 表格
        if stripped.startswith('|'):
            rows, next_i = parse_table(lines, i)
            add_table(doc, rows)
            i = next_i
            continue

        # 按章节处理普通行
        if section == 'refs':
            add_reference(doc, stripped)
        elif section == 'figures':
            add_body(doc, stripped, size=11, line_spacing=1.15, space_after=6)
        else:
            add_body(doc, stripped)
        i += 1

    doc.save(OUT)
    print("已保存:", OUT)


if __name__ == '__main__':
    main()
